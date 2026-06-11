import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Use emerald green colors to match the theme
TITLE_COLOR = RGBColor(4, 120, 87) # Tailwind emerald-700
TEXT_COLOR = RGBColor(55, 65, 81) # Tailwind gray-700

def add_slide(prs, title, content_lines, is_title_slide=False):
    if is_title_slide:
        slide_layout = prs.slide_layouts[0] # Title slide
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title_shape.text = title
        title_shape.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
        subtitle.text = content_lines[0]
    else:
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        body_shape = slide.placeholders[1]
        
        title_shape.text = title
        title_shape.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
        
        tf = body_shape.text_frame
        tf.clear()
        
        for idx, line in enumerate(content_lines):
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(20)
            p.font.color.rgb = TEXT_COLOR
            if idx == 0:
                p.level = 0
            else:
                p.level = 0

# Slide 1
add_slide(prs, "EcoTrack", ["Application de suivi de consommation énergétique\nPrésentation de projet"], is_title_slide=True)

# Slide 2
add_slide(prs, "Problématique", [
    "• Les ménages manquent souvent de visibilité en temps réel sur leur consommation.",
    "• La facture annuelle/mensuelle arrive trop tard pour ajuster ses habitudes.",
    "• Enjeu financier et environnemental majeur (inflation, transition écologique).",
    "• Les solutions classiques sont peu intuitives et non géolocalisées."
])

# Slide 3
add_slide(prs, "Solution Proposée", [
    "• Plateforme moderne et centralisée pour renseigner et suivre les relevés.",
    "• Interface unifiée, moderne et apaisante (Thème 'Émeraude').",
    "• Localisation fine (ville, quartier) pour des analyses comparatives.",
    "• Navigation fluide sans rechargements (React - Single Page Application)."
])

# Slide 4
add_slide(prs, "Technologies Utilisées", [
    "• Backend: Laravel (PHP) - API REST, Sécurité, Modélisation (Eloquent).",
    "• Frontend: React (JavaScript) - Interface dynamique, composants réutilisables.",
    "• UI/UX: TailwindCSS - Design system cohérent et moderne.",
    "• Base de données: MySQL 8.0 - Gestion des relations hiérarchiques.",
    "• Déploiement: Docker & Docker-Compose - Isolation et reproductibilité."
])

# Slide 5
add_slide(prs, "Architecture Conteneurisée (Docker)", [
    "• Architecture Client-Serveur découplée.",
    "• 3 conteneurs isolés et communicants :",
    "    1. Frontend (React) sur port 3000",
    "    2. Backend API (Laravel) sur port 8000",
    "    3. Base de données (MySQL) isolée",
    "• Communication via requêtes HTTP JSON (Axios).",
    "• Exécution automatique des migrations au lancement."
])

# Slide 6
add_slide(prs, "Fonctionnalités Principales", [
    "• Authentification avancée : Inscription dynamique et imbriquée (Ville -> Quartier).",
    "• Tableau de Bord : Vue d'ensemble des statistiques de l'utilisateur.",
    "• CRUD des relevés : Ajout, modification, suppression et suivi des compteurs.",
    "• Performance : Gestion des états globaux côté client (Redux) pour réduire les appels API."
])

# Slide 7
add_slide(prs, "Base de Données & Modélisation", [
    "• Tables principales : Users, Cities, Quartiers, Readings.",
    "• Relations hiérarchiques pensées pour l'analyse :",
    "    - 1 Ville contient plusieurs Quartiers",
    "    - 1 Quartier abrite plusieurs Utilisateurs",
    "    - 1 Utilisateur possède plusieurs Relevés",
    "• Structure préparée pour l'agrégation de données par zone géographique."
])

# Slide 8
add_slide(prs, "Difficultés Techniques Résolues", [
    "• Chargement asynchrone : Gestion des listes déroulantes de quartiers après sélection de la ville.",
    "• Conflits d'états React : Nettoyage rigoureux des erreurs au changement de page.",
    "• Problèmes d'environnement : Les erreurs de commandes Artisan résolues grâce à l'exécution au sein des conteneurs Docker."
])

# Slide 9
add_slide(prs, "Améliorations Futures", [
    "• Intégration IoT : Connexion directe aux compteurs communicants (ex: Linky).",
    "• Gamification : Comparaison anonymisée des consommations moyennes par quartier.",
    "• Alertes intelligentes : Notifications automatiques en cas de pic anormal."
])

# Slide 10
add_slide(prs, "Conclusion", [
    "• Réalisation d'une application utile, esthétique et techniquement robuste.",
    "• Mise en pratique de concepts avancés : Docker, découplage, gestion d'états.",
    "• Préparation aux environnements de production et au travail d'équipe.",
    "",
    "Merci de votre attention ! Avez-vous des questions ?"
])

prs.save('/home/amine/energy-app/Presentation_Soutenance_EcoTrack.pptx')
print("Presentation generated successfully at /home/amine/energy-app/Presentation_Soutenance_EcoTrack.pptx")
